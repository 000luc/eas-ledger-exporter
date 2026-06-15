from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from eas_ledger_exporter.errors import ValidationError
from eas_ledger_exporter.validation import validate_export


REQUIRED_HEADERS = (
    "公司",
    "记账日期",
    "期间",
    "凭证类型",
    "凭证号",
    "状态",
    "制单",
    "审核",
    "出纳",
    "过账",
    "摘要",
    "科目编码",
    "科目名称",
    "币别",
    "核算项目",
    "原币金额",
    "借方",
    "贷方",
    "业务日期",
    "来源系统",
    "审核退回",
    "来源类型",
    "参考信息",
    "附件",
    "计量单位",
    "数量",
    "制单人",
    "结算号",
)


def _write_export(path: Path, headers: tuple[str, ...], rows: list[tuple]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("凭证序时簿")
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _full_row(company: str = "广电计量检测集团股份有限公司", period: str = "2026.6") -> tuple:
    return (
        company,
        "2026-06-01",
        period,
        "记",
        "001",
        "已过账",
        "制单人",
        "审核人",
        "出纳人",
        "已过账",
        "摘要",
        "1001",
        "库存现金",
        "人民币",
        None,
        None,
        100.0,
        None,
        "2026-06-01",
        "EAS",
        None,
        None,
        None,
        None,
        None,
        None,
        "制单人",
        None,
    )


@pytest.fixture
def valid_export(tmp_path):
    path = tmp_path / "001.公司_202606_凭证序时簿.xlsx"
    _write_export(path, REQUIRED_HEADERS, [_full_row(), _full_row()])
    return path


@pytest.fixture
def empty_export(tmp_path):
    path = tmp_path / "001.公司_202606_凭证序时簿.xlsx"
    _write_export(path, REQUIRED_HEADERS, [])
    return path


@pytest.fixture
def wrong_period_export(tmp_path):
    path = tmp_path / "001.公司_202606_凭证序时簿.xlsx"
    _write_export(path, REQUIRED_HEADERS, [_full_row(period="2026.5")])
    return path


@pytest.fixture
def wrong_company_export(tmp_path):
    path = tmp_path / "001.公司_202606_凭证序时簿.xlsx"
    _write_export(path, REQUIRED_HEADERS, [_full_row(company="其他公司")])
    return path


@pytest.fixture
def missing_header_export(tmp_path):
    path = tmp_path / "001.公司_202606_凭证序时簿.xlsx"
    headers = tuple(h for h in REQUIRED_HEADERS if h != "借方")
    _write_export(path, headers, [_full_row()])
    return path


def test_validate_accepts_matching_company_and_period(valid_export):
    result = validate_export(valid_export, "广电计量检测集团股份有限公司", 2026, 6)
    assert result.row_count == 2
    assert result.empty is False


def test_validate_accepts_empty_data(empty_export):
    result = validate_export(empty_export, "广电计量检测集团股份有限公司", 2026, 6)
    assert result.row_count == 0
    assert result.empty is True


def test_validate_rejects_wrong_period(wrong_period_export):
    with pytest.raises(ValidationError, match="期间不一致"):
        validate_export(wrong_period_export, "广电计量检测集团股份有限公司", 2026, 6)


def test_validate_rejects_wrong_company(wrong_company_export):
    with pytest.raises(ValidationError, match="公司不一致"):
        validate_export(wrong_company_export, "广电计量检测集团股份有限公司", 2026, 6)


def test_validate_rejects_missing_header(missing_header_export):
    with pytest.raises(ValidationError, match="缺少表头"):
        validate_export(missing_header_export, "广电计量检测集团股份有限公司", 2026, 6)


def test_validate_accepts_period_with_leading_zero(valid_export):
    path = valid_export
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("凭证序时簿")
    sheet.append(REQUIRED_HEADERS)
    row = list(_full_row())
    row[2] = "2026.06"
    sheet.append(row)
    workbook.save(path)
    workbook.close()

    result = validate_export(path, "广电计量检测集团股份有限公司", 2026, 6)
    assert result.row_count == 1


def test_validate_skips_blank_rows(valid_export):
    path = valid_export
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("凭证序时簿")
    sheet.append(REQUIRED_HEADERS)
    sheet.append(_full_row())
    sheet.append((None,) * len(REQUIRED_HEADERS))
    sheet.append(_full_row())
    workbook.save(path)
    workbook.close()

    result = validate_export(path, "广电计量检测集团股份有限公司", 2026, 6)
    assert result.row_count == 2


def test_validate_rejects_invalid_period_format(valid_export):
    path = valid_export
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("凭证序时簿")
    sheet.append(REQUIRED_HEADERS)
    row = list(_full_row())
    row[2] = "2026/6"
    sheet.append(row)
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValidationError, match="期间格式无效"):
        validate_export(path, "广电计量检测集团股份有限公司", 2026, 6)
