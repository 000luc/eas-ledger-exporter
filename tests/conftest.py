from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook


@pytest.fixture
def make_config_book(tmp_path):
    def make(
        *,
        year=2026,
        month=6,
        output_dir=None,
        companies=(("001", " 广州本部 ", " 是 "), ("002", "深圳公司", "否")),
        period_sheet: bool = True,
        companies_sheet: bool = True,
        filename: str = "info.xlsx",
    ) -> Path:
        path = tmp_path / filename
        workbook = Workbook()
        workbook.remove(workbook.active)

        if period_sheet:
            period = workbook.create_sheet("执行期间")
            period.append(("年份", "月份", "输出路径"))
            period.append(
                (
                    year,
                    month,
                    str(output_dir if output_dir is not None else tmp_path / "序时账"),
                )
            )

        if companies_sheet:
            company_sheet = workbook.create_sheet("执行操作的公司")
            company_sheet.append(("公司编号", "公司名称", "是否执行"))
            for row in companies:
                company_sheet.append(row)

        workbook.save(path)
        workbook.close()
        return path

    return make
