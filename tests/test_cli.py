from __future__ import annotations

from pathlib import Path

import pytest

from eas_ledger_exporter.cli import build_parser, main


def test_parser_requires_config():
    parser = build_parser()
    with pytest.raises(SystemExit):
        parser.parse_args([])


def test_parser_accepts_limit_and_dry_run(tmp_path):
    parser = build_parser()
    args = parser.parse_args(
        ["--config", str(tmp_path / "info.xlsx"), "--limit", "3", "--dry-run"]
    )
    assert args.config == tmp_path / "info.xlsx"
    assert args.limit == 3
    assert args.dry_run is True


def test_main_rejects_missing_config(tmp_path):
    config_path = tmp_path / "missing.xlsx"
    assert main(["--config", str(config_path)]) == 1


def test_main_dry_run_lists_companies(tmp_path, capsys):
    from openpyxl import Workbook

    config_path = tmp_path / "info.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    period = workbook.create_sheet("执行期间")
    period.append(("年份", "月份", "输出路径"))
    period.append((2026, 6, str(tmp_path / "out")))
    companies = workbook.create_sheet("执行操作的公司")
    companies.append(("公司编号", "公司名称", "是否执行"))
    companies.append(("001", "广州本部", "是"))
    companies.append(("002", "深圳公司", "是"))
    workbook.save(config_path)
    workbook.close()

    assert main(["--config", str(config_path), "--dry-run"]) == 0
    captured = capsys.readouterr()
    assert "干运行模式" in captured.out
    assert "001" in captured.out
    assert "002" in captured.out


def test_main_dry_run_respects_limit(tmp_path, capsys):
    from openpyxl import Workbook

    config_path = tmp_path / "info.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    period = workbook.create_sheet("执行期间")
    period.append(("年份", "月份", "输出路径"))
    period.append((2026, 6, str(tmp_path / "out")))
    companies = workbook.create_sheet("执行操作的公司")
    companies.append(("公司编号", "公司名称", "是否执行"))
    companies.append(("001", "广州本部", "是"))
    companies.append(("002", "深圳公司", "是"))
    workbook.save(config_path)
    workbook.close()

    assert main(["--config", str(config_path), "--limit", "1", "--dry-run"]) == 0
    captured = capsys.readouterr()
    assert "001" in captured.out
    assert "002" not in captured.out
