from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .errors import ValidationError


REQUIRED_HEADERS = (
    "公司",
    "记账日期",
    "期间",
    "凭证类型",
    "凭证号",
    "状态",
    "制单",
    "审核",
    "出纳",
    "过账",
    "摘要",
    "科目编码",
    "科目名称",
    "币别",
    "核算项目",
    "原币金额",
    "借方",
    "贷方",
    "业务日期",
    "来源系统",
    "审核退回",
    "来源类型",
    "参考信息",
    "附件",
    "计量单位",
    "数量",
    "制单人",
    "结算号",
)


@dataclass(frozen=True)
class ValidationResult:
    row_count: int
    empty: bool


def _normalize_period(value: object) -> tuple[int, int]:
    text = str(value).strip()
    if "." not in text:
        raise ValidationError(f"期间格式无效：{text!r}")
    year_text, month_text = text.split(".", 1)
    try:
        return int(year_text), int(month_text)
    except ValueError as exc:
        raise ValidationError(f"期间格式无效：{text!r}") from exc


def validate_export(path: Path, company: str, year: int, month: int) -> ValidationResult:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(f"无法打开导出文件：{path}") from exc

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = tuple(next(rows))
        except StopIteration:
            raise ValidationError("工作表没有表头")

        missing = [name for name in REQUIRED_HEADERS if name not in headers]
        if missing:
            raise ValidationError(f"缺少表头：{missing}")

        company_index = headers.index("公司")
        period_index = headers.index("期间")

        row_count = 0
        for row in rows:
            if not any(value is not None for value in row):
                continue
            row_count += 1
            company_value = row[company_index]
            if company_value is not None and str(company_value).strip() != company:
                raise ValidationError(
                    f"公司不一致：期望 {company!r}，实际 {company_value!r}"
                )
            period_value = row[period_index]
            if period_value is not None:
                period = _normalize_period(period_value)
                if period != (year, month):
                    raise ValidationError(
                        f"期间不一致：期望 {year}.{month}，实际 {period_value!r}"
                    )
    finally:
        workbook.close()

    return ValidationResult(row_count=row_count, empty=row_count == 0)
