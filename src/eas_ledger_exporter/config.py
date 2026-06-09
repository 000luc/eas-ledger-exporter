from __future__ import annotations

from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .errors import ConfigError
from .models import Company, ExportConfig


REQUIRED_SHEETS = ("执行期间", "执行操作的公司")
CORRUPT_FILE_ERRORS = (BadZipFile, InvalidFileException, ParseError, EOFError)
OPEN_WORKBOOK_ERRORS = CORRUPT_FILE_ERRORS + (KeyError,)


def _required_integer(value: object, label: str) -> int:
    if isinstance(value, bool):
        raise ConfigError(f"{label}必须是整数")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdecimal():
        return int(value.strip())
    raise ConfigError(f"{label}必须是整数")


def _company_code(value: object, row_number: int) -> str:
    if isinstance(value, bool) or value is None:
        raise ConfigError(f"第{row_number}行公司编号不能为空且必须是整数")

    if isinstance(value, int):
        number = value
    elif isinstance(value, float) and value.is_integer():
        number = int(value)
    elif isinstance(value, str):
        text = value.strip()
        if not text or not text.isdecimal():
            raise ConfigError(f"第{row_number}行公司编号必须是整数")
        number = int(text)
    else:
        raise ConfigError(f"第{row_number}行公司编号必须是整数")

    if not 0 <= number <= 999:
        raise ConfigError(f"第{row_number}行公司编号必须在000至999之间")
    return f"{number:03d}"


def _company_name(value: object, row_number: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ConfigError(f"第{row_number}行公司名称不能为空")
    return value.strip()


def _parse_workbook(workbook) -> ExportConfig:
    for sheet_name in REQUIRED_SHEETS:
        if sheet_name not in workbook.sheetnames:
            raise ConfigError(f"缺少工作表“{sheet_name}”")

    period = workbook["执行期间"]
    year = _required_integer(period["A2"].value, "年份")
    month = _required_integer(period["B2"].value, "月份")
    output = period["C2"].value

    if not 2000 <= year <= 2100:
        raise ConfigError("年份必须在2000至2100之间")
    if not 1 <= month <= 12:
        raise ConfigError("月份必须是1至12")
    if not isinstance(output, str) or not output.strip():
        raise ConfigError("输出路径不能为空")

    companies: list[Company] = []
    seen_codes: set[str] = set()
    rows = workbook["执行操作的公司"].iter_rows(
        min_row=2, max_col=3, values_only=True
    )
    for row_number, (code, name, enabled) in enumerate(rows, start=2):
        if not isinstance(enabled, str) or enabled.strip() != "是":
            continue
        normalized_code = _company_code(code, row_number)
        if normalized_code in seen_codes:
            raise ConfigError(f"公司编号重复：{normalized_code}")
        seen_codes.add(normalized_code)
        companies.append(
            Company(normalized_code, _company_name(name, row_number))
        )

    if not companies:
        raise ConfigError("没有启用的公司")

    return ExportConfig(
        year=year,
        month=month,
        output_dir=Path(output.strip()),
        companies=tuple(companies),
    )


def load_config(path: Path) -> ExportConfig:
    path = Path(path)
    if not path.is_file():
        raise ConfigError(f"配置文件不存在：{path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except OPEN_WORKBOOK_ERRORS as exc:
        raise ConfigError(f"配置文件损坏或格式无效：{path}") from exc
    except OSError as exc:
        raise ConfigError(f"无法读取配置文件“{path}”：{exc}") from exc

    try:
        try:
            return _parse_workbook(workbook)
        except CORRUPT_FILE_ERRORS as exc:
            raise ConfigError(f"配置文件损坏或格式无效：{path}") from exc
    finally:
        workbook.close()
